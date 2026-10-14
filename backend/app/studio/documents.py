"""Read an uploaded resume or template file, and report on its layout.

The text comes from the same extractor the job matcher uses. The report
adds what only the file can tell: page count, tables, a second text
column, contact details tucked into a page header, and, for a DOCX
template, which paragraphs are section headings and which styles they
use, so a tailored version can be written back in the owner's design.
"""

from __future__ import annotations

import io

from backend.app.matching.parsing import ResumeParseError, extract_resume_text
from backend.app.studio.sections import EMAIL, heading_kind


MEDIA_TYPES = {
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".txt": "text/plain",
    ".md": "text/plain",
}


def media_type(filename: str) -> str | None:
    lowered = filename.lower()
    return next((t for ext, t in MEDIA_TYPES.items() if lowered.endswith(ext)), None)


def _pdf_report(payload: bytes) -> dict:
    import pdfplumber

    report: dict = {"pages": 0, "columns": False}
    split_lines = total_lines = 0
    with pdfplumber.open(io.BytesIO(payload)) as pdf:
        report["pages"] = len(pdf.pages)
        for page in pdf.pages[:3]:
            words = page.extract_words()
            rows: dict[int, list] = {}
            for w in words:
                rows.setdefault(round(w["top"] / 3), []).append(w)
            for row in rows.values():
                row.sort(key=lambda w: w["x0"])
                total_lines += 1
                gaps = [b["x0"] - a["x1"] for a, b in zip(row, row[1:])]
                # A wide gap in the middle third of the page, on many
                # lines, is two columns rather than a right-aligned date.
                if any(g > page.width * 0.12 and page.width * 0.3 < b["x0"] < page.width * 0.7
                       for g, b in zip(gaps, row[1:])):
                    split_lines += 1
    report["columns"] = total_lines >= 10 and split_lines / total_lines > 0.35
    return report


def _docx_report(payload: bytes) -> dict:
    import docx

    document = docx.Document(io.BytesIO(payload))
    headings = []
    for index, paragraph in enumerate(document.paragraphs):
        kind = heading_kind(paragraph.text)
        if kind:
            headings.append({"index": index, "kind": kind, "text": paragraph.text.strip(),
                             "style": paragraph.style.name if paragraph.style is not None else ""})
    header_text = " ".join(
        p.text for section in document.sections for part in (section.header, section.footer) for p in part.paragraphs
    )
    return {
        "pages": None,
        "tables": len(document.tables),
        "headings": headings,
        "header_footer_contact": bool(EMAIL.search(header_text)),
        "paragraphs": len(document.paragraphs),
    }


def read_document(*, payload: bytes, filename: str) -> tuple[str, str, dict]:
    """Text, parser name and layout report. An unreadable file returns
    empty text and says why, so it can still be stored and inspected."""

    lowered = filename.lower()
    parser = "pdfplumber" if lowered.endswith(".pdf") else "python-docx" if lowered.endswith(".docx") else "text"
    report: dict = {}
    try:
        if parser == "pdfplumber":
            report = _pdf_report(payload)
        elif parser == "python-docx":
            report = _docx_report(payload)
    except Exception as error:  # a damaged file is reported, not fatal
        report = {"error": f"The file could not be opened: {type(error).__name__}."}
    try:
        text = extract_resume_text(payload=payload, filename=filename)
    except ResumeParseError as error:
        report["error"] = str(error)
        text = ""
    report["characters"] = len(text)
    return text, parser, report
