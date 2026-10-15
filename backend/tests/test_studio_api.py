"""Resume Studio end to end over HTTP, on SQLite, with fictional files."""

import io
from datetime import datetime, timezone

import docx
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import create_engine, select
from sqlalchemy.orm import Session, sessionmaker
from sqlalchemy.pool import StaticPool

from backend.app.api.deps import get_session
from backend.app.api.main import create_app
from backend.app.auth.guard import require_auth
from backend.app.config import get_settings
from backend.app.db.base import Base
from backend.app.db.models import (
    CareerFactRecord,
    JobEvaluationRecord,
    JobRecord,
    JobResumeScoreRecord,
    ResumeRecord,
)
from backend.tests.test_studio_ats import JOB, RESUME


NOW = datetime(2026, 10, 14, 12, tzinfo=timezone.utc)


def _docx(text: str, *, table: bool = False) -> bytes:
    document = docx.Document()
    for line in text.splitlines():
        document.add_paragraph(line)
    if table:
        document.add_table(rows=1, cols=2)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


@pytest.fixture(name="factory")
def fixture_factory():
    engine = create_engine("sqlite+pysqlite:///:memory:", connect_args={"check_same_thread": False},
                           poolclass=StaticPool)
    Base.metadata.create_all(engine)
    return sessionmaker(bind=engine, class_=Session, expire_on_commit=False)


@pytest.fixture(name="client")
def fixture_client(factory):
    get_settings.cache_clear()

    def override():
        with factory() as session:
            yield session

    app = create_app()
    app.dependency_overrides[get_session] = override
    app.dependency_overrides[require_auth] = lambda: None
    with factory.begin() as session:
        job = JobRecord(source="greenhouse", source_account="pinecrest", external_id="1",
                        company="Pinecrest AI", title="Software Engineer", location="Seattle, WA",
                        description=JOB, official_url="https://boards.example.com/pinecrest/1",
                        content_hash="h", first_seen_at=NOW, last_seen_at=NOW, is_active=True)
        session.add(job)
        session.flush()
        session.add(JobEvaluationRecord(job_id=job.id, eligibility_status="PASS",
                                        role_family="SOFTWARE_ENGINEERING", role_priority="PRIMARY",
                                        rule_version="test", reason_codes=["NO_HARD_BLOCKER"],
                                        reasons=["No hard eligibility blocker detected."],
                                        content_hash="h", requirements_verified=True, evaluated_at=NOW))
    yield TestClient(app)
    get_settings.cache_clear()


def _upload(client, data: bytes, name="alex-rivera.docx", kind="resume"):
    return client.post("/api/studio/documents", files={"file": (name, data)}, data={"kind": kind})


def test_uploading_a_resume_suggests_facts_and_a_primary_master(client, factory):
    response = _upload(client, _docx(RESUME))
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["created"] and body["document"]["parser"] == "python-docx"

    facts = client.get("/api/studio/facts").json()["items"]
    assert facts and all(f["status"] == "suggested" and f["origin"] == "parsed" for f in facts)
    assert any(f["kind"] == "experience" and f["organization"] == "Brightmoor Health" for f in facts)

    overview = client.get("/api/studio").json()
    assert overview["primary"]["kind"] == "master" and overview["primary"]["is_primary"]
    # The matcher now scores against it.
    with factory() as session:
        assert session.scalar(select(ResumeRecord).where(ResumeRecord.is_active.is_(True))) is not None


def test_the_same_file_is_stored_once(client):
    data = _docx(RESUME)
    assert _upload(client, data).json()["created"] is True
    assert _upload(client, data).json()["created"] is False
    assert len(client.get("/api/studio").json()["documents"]) == 1


def test_the_original_file_comes_back_byte_for_byte(client):
    data = _docx(RESUME)
    document = _upload(client, data).json()["document"]
    assert client.get(f"/api/studio/documents/{document['id']}/file").content == data


def test_the_document_view_shows_what_was_read(client):
    document = _upload(client, _docx(RESUME, table=True)).json()["document"]
    detail = client.get(f"/api/studio/documents/{document['id']}").json()
    assert detail["report"]["tables"] == 1
    assert [s["kind"] for s in detail["structure"]["sections"]] == ["experience", "projects", "education", "skills"]
    assert detail["structure"]["emails"] == ["alex.rivera@example.com"]


def test_bad_files_are_refused_with_a_reason(client):
    assert _upload(client, b"", "empty.pdf").status_code == 400
    assert _upload(client, b"x" * 10, "resume.exe").status_code == 400
    assert _upload(client, b"plain text", "template.txt", kind="template").status_code == 400


def test_facts_change_with_an_audit_trail(client):
    _upload(client, _docx(RESUME))
    fact = next(f for f in client.get("/api/studio/facts").json()["items"] if f["kind"] == "experience")

    edited = client.patch(f"/api/studio/facts/{fact['id']}", json={"title": "Software Engineer II", "status": "verified"})
    assert edited.json()["status"] == "verified" and edited.json()["title"] == "Software Engineer II"

    history = client.get(f"/api/studio/facts/{fact['id']}/history").json()["items"]
    assert [e["action"] for e in history] == ["created", "edited", "status"]
    assert history[1]["before"]["title"] == "Software Engineer"

    assert client.patch(f"/api/studio/facts/{fact['id']}", json={"status": "invented"}).status_code == 400


def test_a_typed_fact_is_verified_and_nothing_is_ever_hard_deleted(client, factory):
    made = client.post("/api/studio/facts", json={"kind": "project", "title": "Trailhead Maps",
                                                   "skills": ["TypeScript", "react"]}).json()
    assert made["status"] == "verified" and made["skills"] == ["react", "typescript"]
    client.patch(f"/api/studio/facts/{made['id']}", json={"status": "deprecated"})
    with factory() as session:
        assert session.get(CareerFactRecord, made["id"]).status == "deprecated"


def test_a_corrected_version_can_become_primary_and_rescores(client, factory):
    _upload(client, _docx(RESUME))
    primary = client.get("/api/studio").json()["primary"]
    corrected = client.post("/api/studio/versions", json={
        "parent_id": primary["id"], "body_text": RESUME.replace("Granite Ledger", "Granite Ledger Inc")}).json()
    assert corrected["parent_id"] == primary["id"] and corrected["status"] == "approved"

    assert client.post(f"/api/studio/versions/{corrected['id']}/primary").status_code == 200
    overview = client.get("/api/studio").json()
    assert overview["primary"]["id"] == corrected["id"]
    with factory() as session:
        active = session.scalar(select(ResumeRecord).where(ResumeRecord.is_active.is_(True)))
        assert "Granite Ledger Inc" in active.raw_text
        assert session.scalar(select(JobResumeScoreRecord).where(JobResumeScoreRecord.resume_id == active.id))


def test_checking_against_a_stored_job_and_a_pasted_one(client, factory):
    _upload(client, _docx(RESUME))
    version = client.get("/api/studio").json()["primary"]["id"]
    with factory() as session:
        job_id = session.scalar(select(JobRecord.id))

    stored = client.post("/api/studio/check", json={"version_id": version, "job_id": job_id}).json()
    assert stored["state"] == "scored" and stored["job_company"] == "Pinecrest AI"
    assert stored["result"]["evidence"]

    pasted = client.post("/api/studio/check", json={"version_id": version, "job_text": JOB,
                                                     "job_title": "Backend Engineer"}).json()
    assert pasted["job_source"] == "pasted" and pasted["aggregate"] == stored["aggregate"]

    thin = client.post("/api/studio/check", json={"version_id": version, "job_text": "Apply now."}).json()
    assert thin["state"] == "insufficient_job_data" and thin["aggregate"] is None

    history = client.get(f"/api/studio/assessments?version_id={version}").json()["items"]
    assert len(history) == 3
    assert client.get(f"/api/studio/assessments?job_id={job_id}").json()["items"][0]["id"] == stored["id"]


def test_checking_needs_a_job_or_text(client):
    _upload(client, _docx(RESUME))
    version = client.get("/api/studio").json()["primary"]["id"]
    assert client.post("/api/studio/check", json={"version_id": version}).status_code == 400


def test_a_verified_knowledge_base_flags_unbacked_claims(client):
    _upload(client, _docx(RESUME))
    client.post("/api/studio/facts", json={"kind": "skill", "title": "python", "skills": ["python"]})
    version = client.get("/api/studio").json()["primary"]["id"]
    result = client.post("/api/studio/check", json={"version_id": version, "job_text": JOB}).json()["result"]
    assert any("not in your knowledge base" in f["message"] for f in result["findings"])


def test_the_company_directory_counts_open_roles_and_applications(client, factory):
    with factory.begin() as session:
        job = session.scalar(select(JobRecord))
        from backend.app.db.models import JobMarkRecord
        session.add(JobMarkRecord(job_id=job.id, is_saved=False, applied_at=NOW, application_status="applied", updated_at=NOW))
    items = client.get("/api/companies").json()["items"]
    assert items == [{"name": "Pinecrest AI", "tier": items[0]["tier"], "open_roles": 1, "applied": 1,
                      "newest_role_at": items[0]["newest_role_at"]}]


def test_the_dashboard_groups_open_roles_by_tier_and_score(client, factory):
    body = client.get("/api/dashboard").json()
    assert body["total_open"] == 1
    assert sum(t["open"] for t in body["tiers"]) == 1
    assert sum(body["scores"].values()) == 1
    assert len(body["daily"]) == 14
    assert body["top_companies"][0]["name"] == "Pinecrest AI"
